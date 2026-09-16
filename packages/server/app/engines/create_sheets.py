import io
import re

import openpyxl
from openpyxl.styles import Font

from app.schemas.create_sheets import CreateSheetsConfig


def _sanitize_sheet_name(name: str) -> str:
    """Excel sheet 名限制：最多 31 字符，不能包含特定字符。"""
    name = str(name).strip()
    # Excel 不允许的字符
    name = re.sub(r'[\\/*?:\[\]]', "_", name)
    if len(name) > 31:
        name = name[:31]
    if not name:
        name = "Sheet"
    return name


def create_sheets(config: CreateSheetsConfig) -> bytes:
    """根据名称列表批量创建工作表，返回 xlsx 字节。"""
    wb = openpyxl.Workbook()
    # 删除默认 sheet
    wb.remove(wb.active)

    seen: set[str] = set()
    for name in config.names:
        sheet_name = _sanitize_sheet_name(name)
        # 处理重名
        original = sheet_name
        counter = 2
        while sheet_name in seen:
            suffix = f"_{counter}"
            sheet_name = original[: 31 - len(suffix)] + suffix
            counter += 1
        seen.add(sheet_name)

        ws = wb.create_sheet(title=sheet_name)
        if config.headers:
            ws.append(config.headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    buffer.seek(0)
    return buffer.read()

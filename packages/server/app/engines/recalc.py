"""
Excel 公式重算封装。
直接调用 vendored 的 xlsx skill recalc.py 脚本。
"""
import json
import subprocess
import sys
from pathlib import Path


def recalc_excel(file_path: str | Path, timeout: int = 30, force: bool = False) -> dict:
    """
    使用 LibreOffice 重算 Excel 文件中的所有公式。
    返回 JSON 字典，包含 status / total_errors / total_formulas / error_summary。
    """
    file_path = Path(file_path)
    if not file_path.exists():
        return {"error": f"文件不存在: {file_path}"}

    recalc_script = Path(__file__).parent / "_lib" / "recalc.py"
    cmd = [sys.executable, str(recalc_script), str(file_path), str(timeout)]
    if force:
        cmd.append("--force")

    result = subprocess.run(
        cmd,
        capture_output=True,
        text=True,
        timeout=timeout + 10,
    )

    try:
        data = json.loads(result.stdout)
    except json.JSONDecodeError:
        return {
            "error": "无法解析 recalc.py 输出",
            "stdout": result.stdout,
            "stderr": result.stderr,
            "returncode": result.returncode,
        }

    return data


def has_formulas(file_path: str | Path) -> bool:
    """快速检查 Excel 文件是否包含公式。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file_path, data_only=False)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                        wb.close()
                        return True
        wb.close()
    except Exception:
        return False
    return False

import io
import re
from datetime import datetime
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.validate import ValidationConfig, ValidationRule

_ERR_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # 红
_WARN_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # 黄
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")


def _header_style(ws, row: int = 1):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _is_empty(v: Any) -> bool:
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return str(v).strip() == ""


def _to_number(v: Any) -> float | None:
    if _is_empty(v):
        return None
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("，", "")
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def _to_date(v: Any) -> bool:
    if _is_empty(v):
        return False
    if isinstance(v, (datetime, pd.Timestamp, np.datetime64)):
        return True
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%Y年%m月%d日", "%Y-%m-%d %H:%M:%S"):
        try:
            datetime.strptime(s, fmt)
            return True
        except ValueError:
            continue
    return False


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(v, (float, np.floating)):
        f = float(v)
        if f.is_integer():
            return str(int(f))
        return str(round(f, 6))
    return str(v).strip()


def validate_excel(path: str, config: ValidationConfig) -> tuple[bytes, dict[str, Any]]:
    """
    对 Excel 应用校验规则，输出校验报告工作簿。
    返回 (xlsx_bytes, summary)
    """
    if not config.rules:
        raise ValueError("请至少配置一条校验规则")

    wb_src = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if config.source_sheet and config.source_sheet in wb_src.sheetnames:
        sheet_name = config.source_sheet
    else:
        sheet_name = wb_src.sheetnames[0]
    wb_src.close()

    df = pd.read_excel(path, sheet_name=sheet_name, header=config.header_row - 1)
    df.columns = [str(c).strip() for c in df.columns]

    data_start = config.data_start_row or (config.header_row + 1)
    # Excel 行号：df index i -> data_start + i
    def excel_row(i: int) -> int:
        return data_start + int(i)

    # 校验列是否存在
    missing_cols: list[dict] = []
    for rule in config.rules:
        cols = []
        if rule.column:
            cols.append(rule.column)
        if rule.compare_columns:
            cols.extend(rule.compare_columns)
        if rule.baseline_column:
            cols.append(rule.baseline_column)
        for c in cols:
            if c not in df.columns:
                missing_cols.append({
                    "规则编号": rule.rule_id,
                    "规则名称": rule.name,
                    "问题": f"列「{c}」不存在",
                })

    violations: list[dict] = []
    cell_marks: list[tuple[int, str, str]] = []  # (excel_row, column, level)

    def add_violation(rule: ValidationRule, i: int | None, col: str, value: Any, desc: str):
        violations.append({
            "规则编号": rule.rule_id,
            "规则名称": rule.name,
            "等级": "错误" if rule.level == "error" else "警告",
            "行号": excel_row(i) if i is not None else "",
            "列": col,
            "当前值": _fmt(value),
            "问题说明": rule.message or desc,
        })
        if i is not None and col:
            cell_marks.append((excel_row(i), col, rule.level))

    for rule in config.rules:
        col = rule.column

        # ---- required ----
        if rule.rule_type == "required":
            if col not in df.columns:
                continue
            for i, row in df.iterrows():
                if _is_empty(row.get(col)):
                    add_violation(rule, i, col, None, "必填项为空")

        # ---- type ----
        elif rule.rule_type == "type":
            if col not in df.columns:
                continue
            for i, row in df.iterrows():
                v = row.get(col)
                if _is_empty(v):
                    continue
                kind = rule.value_type or "number"
                if kind == "number":
                    if _to_number(v) is None:
                        add_violation(rule, i, col, v, "不是有效数字")
                    elif not rule.allow_negative and _to_number(v) < 0:
                        add_violation(rule, i, col, v, "不允许为负数")
                elif kind == "date":
                    if not _to_date(v):
                        add_violation(rule, i, col, v, "不是有效日期")
                # text 视为始终通过

        # ---- range ----
        elif rule.rule_type == "range":
            if col not in df.columns:
                continue
            for i, row in df.iterrows():
                v = row.get(col)
                if _is_empty(v):
                    continue
                num = _to_number(v)
                if num is None:
                    # 非数值：单列为 type 规则负责，range 不重复报
                    continue
                if not rule.allow_negative and num < 0:
                    add_violation(rule, i, col, v, "不允许为负数")
                    continue
                if rule.min is not None and num < rule.min:
                    add_violation(rule, i, col, v, f"小于下限 {rule.min}")
                if rule.max is not None and num > rule.max:
                    add_violation(rule, i, col, v, f"大于上限 {rule.max}")

        # ---- enum ----
        elif rule.rule_type == "enum":
            if col not in df.columns:
                continue
            allowed = set(str(x).strip() for x in (rule.allowed or []))
            for i, row in df.iterrows():
                v = row.get(col)
                if _is_empty(v):
                    continue
                if str(v).strip() not in allowed:
                    add_violation(rule, i, col, v, f"不在允许值内（{', '.join(sorted(allowed))}）")

        # ---- pattern ----
        elif rule.rule_type == "pattern":
            if col not in df.columns:
                continue
            regex = re.compile(rule.pattern or "")
            for i, row in df.iterrows():
                v = row.get(col)
                if _is_empty(v):
                    continue
                if not regex.search(str(v)):
                    add_violation(rule, i, col, v, f"不匹配格式 {rule.pattern}")

        # ---- compare（列间勾稽）----
        elif rule.rule_type == "compare":
            cols = rule.compare_columns or []
            if len(cols) < 2 or any(c not in df.columns for c in cols):
                continue
            for i, row in df.iterrows():
                nums = [_to_number(row.get(c)) for c in cols]
                if any(n is None for n in nums):
                    continue
                op = rule.compare_op
                target, others = nums[0], nums[1:]
                ok = True
                desc = ""
                if op == "eq":
                    ok = abs(target - others[0]) <= rule.tolerance
                    desc = f"{cols[0]} 应等于 {cols[1]}（容差{rule.tolerance}）"
                elif op == "ne":
                    ok = abs(target - others[0]) > rule.tolerance
                    desc = f"{cols[0]} 不应等于 {cols[1]}"
                elif op == "gt":
                    ok = target > others[0]
                    desc = f"{cols[0]} 应大于 {cols[1]}"
                elif op == "lt":
                    ok = target < others[0]
                    desc = f"{cols[0]} 应小于 {cols[1]}"
                elif op == "gte":
                    ok = target >= others[0]
                    desc = f"{cols[0]} 应大于等于 {cols[1]}"
                elif op == "lte":
                    ok = target <= others[0]
                    desc = f"{cols[0]} 应小于等于 {cols[1]}"
                elif op == "sum_eq":
                    # 支持带符号求和：signs 为其余列各自的符号（默认全 +1）
                    signs = rule.signs or [1.0] * len(others)
                    if len(signs) != len(others):
                        signs = [1.0] * len(others)
                    expected = sum(s * v for s, v in zip(signs, others))
                    ok = abs(target - expected) <= rule.tolerance
                    expr = " ".join(
                        ("+ " if s >= 0 else "- ") + c
                        for s, c in zip(signs, cols[1:])
                    )
                    desc = f"{cols[0]} 应等于 {expr}（容差{rule.tolerance}）"
                if not ok:
                    add_violation(rule, i, cols[0], row.get(cols[0]), desc)

        # ---- threshold ----
        elif rule.rule_type == "threshold":
            if col not in df.columns:
                continue
            for i, row in df.iterrows():
                v = _to_number(row.get(col))
                if v is None:
                    continue
                if rule.threshold_type == "abs":
                    if rule.max is not None and abs(v) > rule.max:
                        add_violation(rule, i, col, v, f"绝对值超过阈值 {rule.max}")
                    if rule.min is not None and abs(v) < rule.min:
                        add_violation(rule, i, col, v, f"绝对值低于阈值 {rule.min}")
                elif rule.threshold_type == "pct_change":
                    base = (
                        _to_number(row.get(rule.baseline_column))
                        if rule.baseline_column
                        else None
                    )
                    # 基数为 0 或本期为 0 时环比无意义，跳过（避免刷屏）
                    if base is None or base == 0 or v == 0:
                        continue
                    pct = (v - base) / abs(base) * 100
                    if rule.max is not None and abs(pct) > rule.max:
                        add_violation(rule, i, col, v, f"环比变化 {pct:.1f}% 超过阈值 {rule.max}%")
                elif rule.threshold_type == "ratio":
                    base = (
                        _to_number(row.get(rule.baseline_column))
                        if rule.baseline_column
                        else None
                    )
                    if base is None or base == 0:
                        continue
                    ratio = v / base * 100
                    if rule.max is not None and ratio > rule.max:
                        add_violation(rule, i, col, v, f"占比 {ratio:.1f}% 超过阈值 {rule.max}%")

    # ---------- 构建报告 ----------
    errors = [v for v in violations if v["等级"] == "错误"]
    warnings = [v for v in violations if v["等级"] == "警告"]

    output = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # 1) 报告汇总
    ws = wb.create_sheet("校验报告")
    rows = [
        ["校验报告", ""],
        ["数据文件", __import__("os").path.basename(path)],
        ["工作表", sheet_name],
        ["表头行", config.header_row],
        ["数据行数", len(df)],
        ["规则数", len(config.rules)],
        ["错误(阻断)", len(errors)],
        ["警告", len(warnings)],
        ["违规合计", len(violations)],
        [
            "结论",
            "未发现阻断性问题"
            if not errors
            else f"存在 {len(errors)} 处阻断性错误，请修正后重试",
        ],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 50
    # 结论着色
    concl = ws.cell(row=len(rows), column=2)
    concl.font = Font(bold=True, color="9C0006" if errors else "006100")

    v_cols = ["规则编号", "规则名称", "等级", "行号", "列", "当前值", "问题说明"]

    # 2) 错误清单
    if errors:
        ws = wb.create_sheet("错误清单")
        ws.append(v_cols)
        _header_style(ws)
        for v in errors:
            ws.append([v.get(c) for c in v_cols])
        for r in range(2, len(errors) + 2):
            for cell in ws[r]:
                cell.fill = _ERR_FILL
        ws.freeze_panes = "A2"
        _auto_width(ws, v_cols)

    # 3) 警告清单
    if warnings:
        ws = wb.create_sheet("警告清单")
        ws.append(v_cols)
        _header_style(ws)
        for v in warnings:
            ws.append([v.get(c) for c in v_cols])
        for r in range(2, len(warnings) + 2):
            for cell in ws[r]:
                cell.fill = _WARN_FILL
        ws.freeze_panes = "A2"
        _auto_width(ws, v_cols)

    # 4) 列结构问题
    if missing_cols:
        ws = wb.create_sheet("列结构问题")
        cols = ["规则编号", "规则名称", "问题"]
        ws.append(cols)
        _header_style(ws)
        for m in missing_cols:
            ws.append([m.get(c) for c in cols])
        for r in range(2, len(missing_cols) + 2):
            for cell in ws[r]:
                cell.fill = _ERR_FILL
        _auto_width(ws, cols)

    # 5) 数据预览（违规单元格标色）
    _write_data_preview(wb, path, sheet_name, df, config, cell_marks)

    if len(wb.sheetnames) == 0:
        wb.create_sheet("校验报告")

    wb.save(output)
    output.seek(0)

    summary = {
        "数据行数": len(df),
        "规则数": len(config.rules),
        "错误": len(errors),
        "警告": len(warnings),
        "违规合计": len(violations),
        "列结构问题": len(missing_cols),
    }
    return output.read(), summary


def _auto_width(ws, headers: list[str]):
    for ci, h in enumerate(headers, start=1):
        width = max(10, min(40, len(str(h)) * 2 + 4))
        ws.column_dimensions[get_column_letter(ci)].width = width


def _write_data_preview(wb, path, sheet_name, df, config, cell_marks):
    """输出数据预览，把违规单元格标上对应等级的底色。"""
    ws = wb.create_sheet("数据预览(违规标色)")
    headers = list(df.columns)
    ws.append(headers)
    _header_style(ws)

    mark_map: dict[tuple[int, str], str] = {}
    for row_no, col, level in cell_marks:
        # 错误优先于警告
        key = (row_no, col)
        if key not in mark_map or level == "error":
            mark_map[key] = level

    data_start = config.data_start_row or (config.header_row + 1)
    preview_n = min(len(df), config.max_preview_rows)
    for i in range(preview_n):
        row = df.iloc[i]
        ws.append([_cell_val(row.get(c)) for c in headers])
        excel_row = data_start + i
        for ci, c in enumerate(headers, start=1):
            lvl = mark_map.get((excel_row, c))
            if lvl == "error":
                ws.cell(row=ws.max_row, column=ci).fill = _ERR_FILL
            elif lvl == "warning":
                ws.cell(row=ws.max_row, column=ci).fill = _WARN_FILL
    ws.freeze_panes = "A2"


def _cell_val(v: Any) -> Any:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v

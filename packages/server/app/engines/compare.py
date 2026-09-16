import io
from datetime import datetime
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

from app.schemas.compare import CompareConfig

# 差异类型颜色
_COLORS = {
    "新增": "C6EFCE",   # 绿
    "删除": "FFC7CE",   # 红
    "修改": "FFEB9C",   # 黄
    "未变": "FFFFFF",
}
_HEADER_FILL = "4472C4"


def _cell_val(v: Any) -> Any:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, datetime):
        return v
    return v


def _read_df(path: str, sheet: str | None, header_row: int) -> tuple[pd.DataFrame, str]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet_name = sheet if sheet and sheet in wb.sheetnames else wb.sheetnames[0]
    wb.close()
    df = pd.read_excel(path, sheet_name=sheet_name, header=header_row - 1)
    df.columns = [str(c).strip() for c in df.columns]
    return df, sheet_name


def _normalize_value(v: Any, config: CompareConfig) -> Any:
    """按配置归一化值用于比较。"""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass

    # 数值比较
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)

    s = str(v)
    if config.ignore_whitespace:
        s = s.strip()
    if config.ignore_case:
        s = s.lower()

    # 尝试把纯数字字符串当数值
    try:
        return float(s)
    except (TypeError, ValueError):
        return s


def _values_equal(a: Any, b: Any, config: CompareConfig) -> bool:
    na = _normalize_value(a, config)
    nb = _normalize_value(b, config)
    if na is None and nb is None:
        return True
    if na is None or nb is None:
        return False
    if isinstance(na, float) and isinstance(nb, float):
        return abs(na - nb) <= config.numeric_tolerance
    return na == nb


def _display(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    # 整数型浮点去掉 .0，避免 40.0 这种显示
    if isinstance(v, (float, np.floating)):
        f = float(v)
        if f.is_integer():
            return str(int(f))
        return str(f)
    return str(v).strip()


def compare_excel(
    old_path: str,
    new_path: str,
    config: CompareConfig,
) -> tuple[bytes, dict[str, Any]]:
    """
    按关键字段匹配两版数据，输出差异对比工作簿。
    返回 (xlsx_bytes, summary)
    """
    old_df, old_sheet = _read_df(
        old_path, config.old_sheet or config.source_sheet, config.header_row
    )
    new_df, new_sheet = _read_df(
        new_path, config.new_sheet or config.source_sheet, config.header_row
    )

    if not config.key_columns:
        raise ValueError("请指定关键字段（key_columns）")

    for k in config.key_columns:
        if k not in old_df.columns:
            cols = ", ".join(old_df.columns)
            raise ValueError(f"关键字段「{k}」在旧版中不存在（旧版列: {cols}）")
        if k not in new_df.columns:
            cols = ", ".join(new_df.columns)
            raise ValueError(f"关键字段「{k}」在新版中不存在（新版列: {cols}）")

    # 决定比对列
    if config.compare_columns:
        compare_cols = [
            c
            for c in config.compare_columns
            if c in old_df.columns and c in new_df.columns
        ]
    else:
        compare_cols = [
            c for c in old_df.columns
            if c not in config.key_columns and c in new_df.columns
        ]
    if not compare_cols:
        raise ValueError("没有可比对的字段（请检查 compare_columns 或两版列名是否一致）")

    # 全部列（用于输出）
    all_cols = list(dict.fromkeys(list(old_df.columns) + list(new_df.columns)))

    def make_key(row) -> tuple:
        """构造匹配键：关键字段的值按配置归一化（忽略空格/大小写）。"""
        parts = []
        for k in config.key_columns:
            s = _display(row.get(k))
            if config.ignore_whitespace:
                s = s.strip()
            if config.ignore_case:
                s = s.lower()
            parts.append(s)
        return tuple(parts)

    def key_label_of(row) -> str:
        """用于展示的键标签：保留原始大小写，仅去空格。"""
        return " | ".join(_display(row.get(k)) for k in config.key_columns)

    # 建立 key -> 行映射（保持顺序，重复 key 取首个并记录）
    old_map: dict[tuple, Any] = {}
    old_dup_keys: list[tuple] = []
    for _, row in old_df.iterrows():
        k = make_key(row)
        if k in old_map:
            old_dup_keys.append(k)
        else:
            old_map[k] = row

    new_map: dict[tuple, Any] = {}
    new_dup_keys: list[tuple] = []
    for _, row in new_df.iterrows():
        k = make_key(row)
        if k in new_map:
            new_dup_keys.append(k)
        else:
            new_map[k] = row

    added: list[dict] = []
    removed: list[dict] = []
    changed: list[dict] = []
    unchanged: list[dict] = []

    # 遍历旧版：找删除 / 修改 / 未变
    for k, old_row in old_map.items():
        key_label = key_label_of(old_row)
        if k not in new_map:
            rec = {"关键字段值": key_label}
            for c in all_cols:
                rec[c] = _cell_val(old_row.get(c))
            removed.append(rec)
        else:
            new_row = new_map[k]
            diffs = []
            for c in compare_cols:
                if not _values_equal(old_row.get(c), new_row.get(c), config):
                    diffs.append((c, _display(old_row.get(c)), _display(new_row.get(c))))
            if diffs:
                for c, ov, nv in diffs:
                    changed.append({
                        "关键字段值": key_label,
                        "变更字段": c,
                        "旧值": ov,
                        "新值": nv,
                    })
            else:
                rec = {"关键字段值": key_label}
                for c in all_cols:
                    rec[c] = _cell_val(new_row.get(c))
                unchanged.append(rec)

    # 遍历新版：找新增
    for k, new_row in new_map.items():
        if k not in old_map:
            key_label = key_label_of(new_row)
            rec = {"关键字段值": key_label}
            for c in all_cols:
                rec[c] = _cell_val(new_row.get(c))
            added.append(rec)

    # ---------- 构建输出 ----------
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # 1) 差异汇总
    summary_rows = [
        ["比对结果", "数量"],
        ["新增记录", len(added)],
        ["删除记录", len(removed)],
        ["修改记录(字段级)", len(changed)],
        ["未变化记录", len(unchanged)],
        ["旧版总行数", len(old_df)],
        ["新版总行数", len(new_df)],
        ["旧版工作表", old_sheet],
        ["新版工作表", new_sheet],
        ["关键字段", ", ".join(config.key_columns)],
        ["比对字段", ", ".join(compare_cols)],
        ["数值容差", config.numeric_tolerance],
        ["忽略空格", "是" if config.ignore_whitespace else "否"],
        ["忽略大小写", "是" if config.ignore_case else "否"],
    ]
    ws = wb.create_sheet("差异汇总")
    for r in summary_rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type="solid")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40

    # 2) 新增记录
    if added:
        _write_list_sheet(wb, "新增记录", added, all_cols, "新增")
    # 3) 删除记录
    if removed:
        _write_list_sheet(wb, "删除记录", removed, all_cols, "删除")
    # 4) 修改记录明细
    if changed:
        cols = ["关键字段值", "变更字段", "旧值", "新值"]
        ws = wb.create_sheet("修改记录")
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type="solid"
            )
        for rec in changed:
            ws.append([rec.get(c) for c in cols])
        for r in range(2, len(changed) + 2):
            for cell in ws[r]:
                cell.fill = PatternFill(
                    start_color=_COLORS["修改"], end_color=_COLORS["修改"], fill_type="solid"
                )
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 20
        ws.freeze_panes = "A2"

    # 5) full 模式：并排全量对比
    if config.output_mode == "full":
        _write_side_by_side(wb, old_df, new_df, config, compare_cols, old_map, new_map, make_key)

    if len(wb.sheetnames) == 0:
        wb.create_sheet("差异汇总")

    wb.save(output)
    output.seek(0)

    summary = {
        "新增": len(added),
        "删除": len(removed),
        "修改字段数": len(changed),
        "未变": len(unchanged),
        "旧版行数": len(old_df),
        "新版行数": len(new_df),
        "旧版重复key": len(old_dup_keys),
        "新版重复key": len(new_dup_keys),
    }
    return output.read(), summary


def _write_list_sheet(wb, title: str, records: list[dict], all_cols: list[str], kind: str):
    # 关键字段值已单独成列，避免与原始列重复
    cols = ["关键字段值"] + [c for c in all_cols if c != "关键字段值"]
    ws = wb.create_sheet(title)
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type="solid")
    for rec in records:
        ws.append([rec.get(c) for c in cols])
    fill = PatternFill(start_color=_COLORS[kind], end_color=_COLORS[kind], fill_type="solid")
    for r in range(2, len(records) + 2):
        for cell in ws[r]:
            cell.fill = fill
    ws.freeze_panes = "A2"


def _write_side_by_side(wb, old_df, new_df, config, compare_cols, old_map, new_map, make_key):
    """并排全量对比：每行显示旧版/新版对应的比对字段值，并标色。"""
    ws = wb.create_sheet("并排对比")
    header = ["关键字段值", "状态"]
    for c in compare_cols:
        header.append(f"旧-{c}")
        header.append(f"新-{c}")
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type="solid")

    # 汇总所有 key，保持新旧顺序
    all_keys: list[tuple] = []
    for k in list(old_map.keys()) + list(new_map.keys()):
        if k not in all_keys:
            all_keys.append(k)

    r = 1
    for k in all_keys:
        old_row = old_map.get(k)
        new_row = new_map.get(k)
        if old_row is not None and new_row is not None:
            diffs = [
                c
                for c in compare_cols
                if not _values_equal(old_row.get(c), new_row.get(c), config)
            ]
            status = "修改" if diffs else "未变"
            if status == "未变" and not config.include_unchanged:
                continue
        elif new_row is not None:
            status = "新增"
        else:
            status = "删除"

        r += 1
        src_row = old_row if old_row is not None else new_row
        key_label = " | ".join(_display(src_row.get(c)) for c in config.key_columns)
        row_vals = [key_label, status]
        for c in compare_cols:
            row_vals.append(_display(old_row.get(c)) if old_row is not None else "")
            row_vals.append(_display(new_row.get(c)) if new_row is not None else "")
        ws.append(row_vals)

        fill = PatternFill(
            start_color=_COLORS[status], end_color=_COLORS[status], fill_type="solid"
        )
        for cell in ws[r]:
            cell.fill = fill
        # 修改的字段单独深色标注
        if status == "修改":
            for ci, c in enumerate(compare_cols):
                if not _values_equal(old_row.get(c), new_row.get(c), config):
                    ws.cell(row=r, column=3 + ci * 2).fill = PatternFill(
                        start_color="FFC000", end_color="FFC000", fill_type="solid")
                    ws.cell(row=r, column=4 + ci * 2).fill = PatternFill(
                        start_color="FFD966", end_color="FFD966", fill_type="solid")

    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 8
